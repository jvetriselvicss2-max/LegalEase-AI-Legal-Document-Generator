from io import BytesIO
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from fpdf import FPDF
from .helpers import sanitize_text

def format_txt(text: str) -> str:
    return sanitize_text(text)

def format_docx(text: str, doc_type: str) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(doc_type.upper())
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    for block in sanitize_text(text).split("\n"):
        if block.strip():
            p = doc.add_paragraph(block.strip())
            p.paragraph_format.space_after = Pt(6)

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Generated with LegalEase")

    output = BytesIO()
    doc.save(output)
    return output.getvalue()

class LegalEasePDF(FPDF):
    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", "I", 8)
        self.cell(0, 10, "Generated with LegalEase", align="C")

def format_pdf(text: str, doc_type: str) -> bytes:
    pdf = LegalEasePDF()
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(0, 10, sanitize_text(doc_type.upper()), align="C")
    pdf.ln(5)
    pdf.set_font("Helvetica", size=11)

    for paragraph in sanitize_text(text).split("\n"):
        if paragraph.strip():
            pdf.multi_cell(0, 7, paragraph.strip())
            pdf.ln(2)

    return bytes(pdf.output())
